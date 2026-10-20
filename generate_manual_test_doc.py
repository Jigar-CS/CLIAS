import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    run.font.name = "Segoe UI"
    if level == 1:
        run.font.size = Pt(17)
        run.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Navy
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        
        # Add bottom accent bar using XML border on heading paragraph if possible
        pPr = h._element.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="3B82F6"/></w:pBdr>')
        pPr.append(pBdr)
    elif level == 2:
        run.font.size = Pt(13.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) # Teal
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
    return h

def add_body_p(doc, text, bold_prefix="", space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.bold = True
        r_bold.font.name = "Segoe UI"
        r_bold.font.size = Pt(10)
        r_bold.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_bullet_p(doc, text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.bold = True
        r_bold.font.name = "Segoe UI"
        r_bold.font.size = Pt(9.5)
        r_bold.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_callout(doc, text, title="IMPORTANT NOTICE", fill_hex="EFF6FF", border_color="2563EB", text_color="1E3A8A"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_title = p.add_run(f"[{title}] ")
    run_title.bold = True
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(9.5)
    run_title.font.color.rgb = RGBColor.from_string(text_color)
    
    run_text = p.add_run(text)
    run_text.font.name = "Segoe UI"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def format_table(table, col_widths, headers, rows_data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], "1E3A8A") # Navy header
        set_cell_margins(hdr_cells[i], top=110, bottom=110, left=110, right=110)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    # Repeat header on subsequent pages
    trPr = table.rows[0]._element.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    # Data rows
    for r_idx, row_values in enumerate(rows_data):
        row = table.add_row()
        # Prevent row split across pages
        r_trPr = row._element.get_or_add_trPr()
        r_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            run = p.add_run(val)
            run.font.name = "Segoe UI"
            run.font.size = Pt(8.8)
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

def create_test_case_table(doc, tc_id, title, role, priority, preconditions, steps, test_data, expected_results):
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="94A3B8")
    
    col_widths = [1.8, 5.0]
    
    # Row 0: Header with TC ID and Title
    hdr_cell = table.rows[0].cells[0]
    hdr_cell.merge(table.rows[0].cells[1])
    set_cell_background(hdr_cell, "1E3A8A")
    set_cell_margins(hdr_cell, top=100, bottom=100, left=140, right=140)
    p = hdr_cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(f"Test Case {tc_id}: ")
    r1.bold = True
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
    r2 = p.add_run(title)
    r2.bold = True
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    meta_rows = [
        ("Role & Priority", f"Role: {role}  |  Priority: {priority}  |  Status: [  ] Pass   [  ] Fail   [  ] Blocked"),
        ("Preconditions", preconditions),
        ("Step-by-Step Flow", steps),
        ("Test Data / Input", test_data),
        ("Expected Result", expected_results)
    ]
    
    for idx, (label, content) in enumerate(meta_rows, start=1):
        c0 = table.rows[idx].cells[0]
        c1 = table.rows[idx].cells[1]
        c0.width = Inches(col_widths[0])
        c1.width = Inches(col_widths[1])
        
        bg_col = "F1F5F9" if idx % 2 == 1 else "F8FAFC"
        set_cell_background(c0, "E2E8F0")
        set_cell_background(c1, bg_col)
        
        set_cell_margins(c0, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r_lbl = p0.add_run(label)
        r_lbl.bold = True
        r_lbl.font.name = "Segoe UI"
        r_lbl.font.size = Pt(8.8)
        r_lbl.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.15
        r_cnt = p1.add_run(content)
        r_cnt.font.name = "Segoe UI"
        r_cnt.font.size = Pt(8.8)
        r_cnt.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)

print("Helper definitions complete.")
